from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw
import json, zipfile

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables' / 'Severino_User_Manual_Icon_Edition.docx'
QA = ROOT / 'tools' / 'manual_qa'
QA.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)
sec.different_first_page_header_footer = True

# Compact reference guide preset; named overrides: editorial cover, navy/gold
# icon-led directory, 10.5-point reference bullets, and linked contents list.
for name, size, color, before, after in [('Normal',11,'20262C',0,6),('Title',32,'173348',0,12),('Subtitle',15,'59656F',0,10),('Heading 1',16,'2E74B5',18,10),('Heading 2',13,'2E74B5',14,7),('Heading 3',12,'1F4D78',10,5)]:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.line_spacing=1.25
    if name.startswith('Heading'): st.font.bold=True; st.paragraph_format.keep_with_next=True
for name,size in [('Manual Step',11),('Directory',9.5),('Small',9),('Contents',11)]:
    st=doc.styles.add_style(name,1); st.base_style=doc.styles['Normal']; st.font.size=Pt(size)
    st.paragraph_format.space_after=Pt(4 if name!='Contents' else 9)
    st.paragraph_format.line_spacing=1.25 if name!='Directory' else 1.1
doc.styles['Small'].font.color.rgb=RGBColor.from_string('59656F')
doc.styles['Manual Step'].paragraph_format.left_indent=Inches(.375)
doc.styles['Manual Step'].paragraph_format.first_line_indent=Inches(-.188)
for level in (1,2,3):
    doc.styles[f'Heading {level}'].font.color.rgb=RGBColor.from_string('173348')
for name,size,before,after in [('Page Entry',14,12,6),('Reference Bullet',10.5,0,5),('Section Kicker',9,0,10)]:
    st=doc.styles.add_style(name,1);st.base_style=doc.styles['Normal'];st.font.size=Pt(size)
    st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.line_spacing=1.15
doc.styles['Page Entry'].base_style=doc.styles['Heading 2']
doc.styles['Page Entry'].font.bold=True
doc.styles['Page Entry'].font.color.rgb=RGBColor.from_string('173348')
doc.styles['Page Entry'].paragraph_format.keep_with_next=True
doc.styles['Section Kicker'].font.color.rgb=RGBColor.from_string('91703D')
doc.styles['Section Kicker'].font.bold=True

header=sec.header.paragraphs[0]; header.style=doc.styles['Small']; header.text='SEVERINO  |  Web Application User Manual'
footer=sec.footer.paragraphs[0]; footer.style=doc.styles['Small']; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Page ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)

def p(text, style=None): return doc.add_paragraph(text,style)
def h(text): return doc.add_heading(text,2)
def note(text):
    par=p('', 'Small'); par.add_run('Note. ').bold=True; par.add_run(text)
def steps(items):
    numbering=doc.part.numbering_part.element
    n=100+len(numbering)
    ab=OxmlElement('w:abstractNum'); ab.set(qn('w:abstractNumId'),str(n))
    lvl=OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'),'0')
    for tag,val in [('start','1'),('numFmt','decimal'),('lvlText','%1.'),('lvlJc','left')]:
        el=OxmlElement('w:'+tag);el.set(qn('w:val'),val);lvl.append(el)
    pp=OxmlElement('w:pPr'); ind=OxmlElement('w:ind');ind.set(qn('w:left'),'540');ind.set(qn('w:hanging'),'270');pp.append(ind)
    tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'num');tab.set(qn('w:pos'),'540');tabs.append(tab);pp.append(tabs);lvl.append(pp);ab.append(lvl);numbering.append(ab)
    num=OxmlElement('w:num');num.set(qn('w:numId'),str(n)); aid=OxmlElement('w:abstractNumId');aid.set(qn('w:val'),str(n));num.append(aid);numbering.append(num)
    for text in items:
        par=p(text,'Manual Step'); np=OxmlElement('w:numPr')
        il=OxmlElement('w:ilvl');il.set(qn('w:val'),'0');ni=OxmlElement('w:numId');ni.set(qn('w:val'),str(n));np.extend([il,ni]);par._p.get_or_add_pPr().append(np)

titles=['Introduction and Access','Customer Page Directory','Administrator Page Directory','Create an Account and Sign In','Browse Products and Manage Favorites','Manage the Cart and Place an Order','Track Orders and Manage Notifications','Update an Account and Submit Feedback','Administrator Access and Store Content','Add and Modify Products','Manage Orders, Users, and Reports','Troubleshooting and Implementation Notes']
def page(title):
    doc.add_page_break(); par=doc.add_heading(title,1)
    ix=titles.index(title)+1
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(ix));b.set(qn('w:name'),f'section_{ix}');par._p.insert(0,b)
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),str(ix));par._p.append(e)

def icon_for(label):
    canvas=Image.new('RGBA',(144,144),(255,255,255,0));d=ImageDraw.Draw(canvas)
    ink='#173348'; gold='#91703D'
    d.rounded_rectangle((2,2,142,142),radius=30,fill='#F2EEE7')
    def line(points,c=ink,w=5):d.line(points,fill=c,width=w,joint='curve')
    def box(coords,r=5):d.rounded_rectangle(coords,radius=r,outline=ink,width=5)
    def circle(coords,c=ink):d.ellipse(coords,outline=c,width=5)
    key={'Home':'home','Shop':'shop','Product Detail':'bottle','Login':'login','Admin Login':'lock','Create Account':'userplus','Verify Email':'mailcheck','Search':'search','Favorites':'heart','Cart':'cart','Checkout':'check','Orders':'truck','Notifications':'bell','Account':'user','Feedback':'message','Billing':'receipt','Email Confirmation':'mailcheck','Dashboard':'grid','Sales Report':'chart','Add Product':'plusbox','Modify Products':'edit','View Users':'users','View Orders':'truck'}.get(label,'info')
    if key=='home':
        line([(29,68),(72,32),(115,68)]);line([(40,61),(40,113),(104,113),(104,61)]);line([(62,113),(62,83),(83,83),(83,113)])
    elif key=='shop':
        box((36,62,109,113));line([(27,63),(37,34),(107,34),(117,63),(27,63)]);line([(57,36),(53,63)]);line([(87,36),(92,63)]);box((65,81,90,113),1)
    elif key=='bottle':
        box((44,59,100,117),10);box((57,29,87,53),2);line([(62,54),(62,59),(82,59),(82,54)]);box((54,77,90,100),2)
    elif key in ('login','lock'):
        if key=='lock':
            box((40,63,104,112));d.arc((52,29,92,87),180,360,fill=ink,width=5);circle((68,79,76,87));line([(72,87),(72,96)],gold)
        else:
            line([(73,34),(108,34),(108,110),(73,110)]);line([(28,72),(84,72)]);line([(68,56),(84,72),(68,88)],gold)
    elif key in ('user','users','userplus'):
        circle((49,29,89,69));d.arc((29,77,109,145),180,360,fill=ink,width=5)
        if key=='userplus':line([(111,45),(111,71)],gold);line([(98,58),(124,58)],gold)
        if key=='users':d.arc((88,35,120,72),250,90,fill=gold,width=4);d.arc((92,80,136,124),260,355,fill=gold,width=4)
    elif key=='mailcheck':
        box((28,42,115,102));line([(29,44),(71,76),(114,44)]);line([(75,107),(86,118),(112,89)],gold,6)
    elif key=='search':
        circle((32,29,90,87));line([(82,80),(114,114)],gold,7)
    elif key=='heart':
        d.arc((30,36,75,81),180,355,fill=ink,width=5);d.arc((70,36,115,81),185,360,fill=ink,width=5);line([(31,58),(33,78),(72,115),(112,78),(114,58)])
    elif key=='cart':
        line([(23,33),(37,33),(52,91),(108,91)]);line([(42,48),(117,48),(107,78),(50,78)]);circle((53,105,63,115));circle((96,105,106,115))
    elif key=='check':
        box((34,28,109,117));line([(52,72),(66,86),(93,57)],gold,6);line([(52,101),(91,101)])
    elif key=='truck':
        box((24,46,79,96),2);line([(79,60),(99,60),(120,80),(120,96),(79,96)]);circle((38,90,60,112));circle((94,90,116,112));line([(93,61),(93,79),(118,79)],gold,4)
    elif key=='bell':
        d.arc((45,36,99,88),180,360,fill=ink,width=5);line([(45,63),(45,88),(34,99),(110,99),(99,88),(99,63)]);d.arc((61,101,83,120),0,180,fill=gold,width=5);line([(72,25),(72,35)])
    elif key=='message':
        line([(30,35),(115,35),(115,95),(61,95),(39,115),(39,95),(30,95),(30,35)]);line([(47,56),(98,56)],gold);line([(47,73),(86,73)])
    elif key=='receipt':
        line([(42,28),(103,28),(103,116),(92,108),(81,116),(70,108),(58,116),(42,108),(42,28)]);line([(55,50),(89,50)],gold);line([(55,70),(89,70)]);line([(55,90),(80,90)])
    elif key=='grid':
        for x,y in [(30,30),(81,30),(30,81),(81,81)]:box((x,y,x+33,y+33))
    elif key=='chart':
        line([(29,28),(29,116),(119,116)]);box((43,82,57,106),1);box((70,61,84,106),1);box((97,38,111,106),1)
    elif key=='plusbox':
        box((30,30,114,114));line([(72,48),(72,96)],gold,6);line([(48,72),(96,72)],gold,6)
    elif key=='edit':
        line([(38,91),(92,37),(109,54),(55,108),(34,114),(38,91)]);line([(82,47),(99,64)],gold);line([(41,91),(55,105)])
    else:
        circle((31,31,113,113));circle((69,47,75,53),gold);line([(72,66),(72,94)],gold,6)
    path=QA/('icon_'+key+'.png');canvas.save(path);return path

def bullet(label,text,keep=False):
    par=p('','Reference Bullet');par.paragraph_format.keep_with_next=keep
    par.paragraph_format.left_indent=Inches(.45);par.paragraph_format.first_line_indent=Inches(-.188)
    np=OxmlElement('w:numPr');il=OxmlElement('w:ilvl');il.set(qn('w:val'),'0');ni=OxmlElement('w:numId');ni.set(qn('w:val'),'90');np.extend([il,ni]);par._p.get_or_add_pPr().append(np)
    par.add_run(label+': ').bold=True;par.add_run(text)

numbering=doc.part.numbering_part.element
ab=OxmlElement('w:abstractNum');ab.set(qn('w:abstractNumId'),'90');lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
for tag,val in [('start','1'),('numFmt','bullet'),('lvlText','•'),('lvlJc','left')]:
    el=OxmlElement('w:'+tag);el.set(qn('w:val'),val);lvl.append(el)
pp=OxmlElement('w:pPr');ind=OxmlElement('w:ind');ind.set(qn('w:left'),'648');ind.set(qn('w:hanging'),'270');pp.append(ind);lvl.append(pp);ab.append(lvl);numbering.append(ab)
num=OxmlElement('w:num');num.set(qn('w:numId'),'90');a=OxmlElement('w:abstractNumId');a.set(qn('w:val'),'90');num.append(a);numbering.append(num)

def icon_entries(rows,widths=None):
    troubleshooting=not rows[0][1].startswith('/')
    area='QUICK HELP' if troubleshooting else ('CUSTOMER PAGES' if rows[0][0]=='Home' else 'ADMINISTRATOR PAGES')
    group_size=3 if troubleshooting else 4
    for i,(name,path,use) in enumerate(rows):
        if i and i%group_size==0:
            doc.add_page_break()
            p(area+'  /  CONTINUED','Section Kicker')
        elif i==0:p(area,'Section Kicker')
        par=p('','Page Entry');r=par.add_run();r.add_picture(str(icon_for(name)),width=Inches(.34))
        r._r.xpath('.//wp:docPr')[0].set('descr',name+' icon')
        par.add_run('  '+name)
        bullet('Recommended Action' if troubleshooting else 'Primary Use',use,True)
        if troubleshooting:bullet('Where',path)
        else:
            parts=path.split('\n');bullet('Page Address',parts[0],True)
            bullet('Access',parts[1] if len(parts)>1 else ('Public sign-in page' if name=='Admin Login' else 'Authorized administrator'))

# Cover: one understated, native geometric perfume icon.
im=Image.new('RGBA',(260,300),(255,255,255,0)); draw=ImageDraw.Draw(im)
draw.rounded_rectangle((62,88,198,260),radius=18,outline='#173348',width=6)
draw.rectangle((92,44,168,87),outline='#173348',width=6)
draw.rectangle((82,140,178,206),outline='#173348',width=4)
im.save(QA/'perfume_icon.png')
par=p('');par.paragraph_format.space_before=Pt(68);par.alignment=1
run=par.add_run();run.add_picture(str(QA/'perfume_icon.png'),width=Inches(.7))
run._r.xpath('.//wp:docPr')[0].set('descr','Minimal outline of a perfume bottle')
par=p('SEVERINO','Title');par.alignment=1
par=p('Perfume Web Application','Subtitle');par.alignment=1
par=p('User Manual and System Documentation','Subtitle');par.alignment=1;par.paragraph_format.space_before=Pt(22)
par=p('Customer and Administrator Guide');par.alignment=1
par=p('Document version 1.1  |  Icon Edition\n6 September 2026','Small');par.alignment=1;par.paragraph_format.space_before=Pt(62)

doc.add_page_break();doc.add_heading('Table of Contents',1)
p('Select a section title to navigate to that section.','Small')
for ix,title in enumerate(titles,1):
    par=p('', 'Contents');link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),f'section_{ix}')
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'173348');rp.append(c);r.append(rp);tx=OxmlElement('w:t');tx.text=title;r.append(tx);link.append(r);par._p.append(link)

page(titles[0])
p('Severino is a perfume shopping web application that enables customers to discover fragrances, manage their accounts, place Cash on Delivery orders, and share feedback. Its administrator area supports product maintenance, order processing, customer records, sales reporting, and storefront content management.')
h('Purpose and scope')
p('This manual explains the purpose of each registered application page and provides practical operating instructions for customers and authorized administrators. It reflects the project source available on 6 September 2026. Page names and button labels are retained where they help users locate controls.')
h('Access requirements')
p('Open the website address supplied by the store or project administrator in a web browser. An internet connection is needed for account services, product data, uploads, and weather information. Registration requires access to an email inbox for a one-time password (OTP).')
h('User roles')
p('Visitor: May view Home, Shop, and Product Detail, and access registration or sign-in pages. Shopping actions that require an account redirect to Login.')
p('Customer: May use Search, Favorites, Cart, Checkout, Orders, Notifications, Account, and Feedback after signing in.')
p('Administrator: Uses Admin Studio to maintain store information and process orders. Administrator accounts are directed to the administrator area when accessing customer pages.')
h('Navigation and notation')
p('Paths in the page directory are appended to the website address. The placeholder :id represents a specific product identifier; open a product card to reach its actual address. On smaller displays, open the navigation menu and scroll tables horizontally when necessary.')
note('This is an operating guide based on source review, not a certification of a live deployment. Display-only features and relevant implementation limits are identified at the end of the manual.')

page(titles[1])
icon_entries([
('Home','/\nPublic','View store content, banners, featured content, and adaptive scent recommendations.'),
('Shop','/shop\nPublic','Browse the collection; filter by name, category, and price; sort results.'),
('Product Detail','/product/:id\nPublic','View images, price, stock, and reviews. Sign in to purchase, save favorites, or submit a review.'),
('Login','/login\nSigned-out users','Sign in using email or mobile number and password, or Google when configured.'),
('Create Account','/create-account\nSigned-out users','Register through profile, email verification, password, and address steps.'),
('Verify Email','/verify-email\nSigned-out users','Complete a separate email challenge when directed to this page.'),
('Search','/search\nCustomer','Find perfumes by name, scent note, and category.'),
('Favorites','/favorites\nCustomer','View saved perfumes and manage favorite selections.'),
('Cart','/cart\nCustomer','Change quantities, remove items, and select items for checkout.'),
('Checkout','/checkout\nCustomer','Complete Shipping, Contact, and Review; submit a COD order.'),
('Orders','/orders\nCustomer','View item status, cancel eligible items, buy again, or open product ratings.'),
('Notifications','/notifications\nCustomer','View order updates; mark messages opened and delete opened entries.'),
('Account','/account\nCustomer','Update profile and address; select a theme; manage an available password.'),
('Feedback','/feedback\nCustomer','Send an order reference, rating, message, and optional attachment.'),
('Billing','/billing\nCustomer','Display-only billing form; use Checkout to submit an order.'),
('Email Confirmation','/email-confirmation\nPublic','Static confirmation preview; displayed details are examples.')])

page(titles[2])
icon_entries([
('Admin Login','/admin/login','Sign in using an authorized administrator account.'),
('Dashboard','/admin','View summary indicators and manage storefront images, banners, pop-up content, and announcements.'),
('Sales Report','/admin/sales','Review sales totals and weekly trends; open a printable PDF report.'),
('Feedback','/admin/feedback','Read customer feedback, order references, ratings, and attachments.'),
('Add Product','/admin/add-product','Create a product with pricing, stock, scent information, images, category, and status.'),
('Modify Products','/admin/products','Select an existing product and edit its listing information.'),
('View Users','/admin/users','Review customer records and remove a user when required.'),
('View Orders','/admin/orders','Review orders and update individual item statuses; remove records when required.')])
h('Administrator navigation')
p('After signing in, use the Admin Studio sidebar to select a page. On a narrow screen, select Menu to reveal the links. Select Sign Out when the administrative task is complete.')
h('Pages outside the active navigation')
p('The repository also contains AdminReviews.jsx and ProductReviews.jsx, but the application router does not register standalone pages for them. Product reviews are available within Product Detail. No separate review-management page is included in this manual’s active page directory.')
note('Administrator pages require both an administrator role and an active administrator session. A customer account cannot use them simply by entering an administrator URL.')

page(titles[3])
h('Register a customer account')
steps(['Open Login and choose the account-creation option, or open /create-account.',
'Enter First Name, Last Name, Email, and Mobile Number. Follow the mobile field’s displayed country-code format.',
'Select Send OTP. Retrieve the six-digit code from the supplied email inbox.',
'Select Continue to Email Verification. Enter the code and select Verify OTP.',
'After verification, enter a password of at least eight characters and repeat it in Confirm Password.',
'Select Continue to Address Info. Complete Barangay, City/Municipality, Province/State, ZIP Code, and Country. Street Address is optional.',
'Submit the account-creation form. Wait for the success confirmation and return to Login.'])
note('Resend OTP becomes available after the displayed cooldown. Changing the email address invalidates the earlier verification and requires verification of the new address.')
h('Sign in and sign out')
steps(['Open /login. Enter the registered email address or contact number and password.',
'Read the terms and conditions, then select the acceptance checkbox.',
'Select Sign In. The application opens the customer area or the appropriate authorized destination.',
'If using Google, accept the terms and select Continue with Google, then complete the provider’s sign-in flow. This option depends on deployment configuration.',
'When finished, use the account navigation’s sign-out option.'])
note('If directed to /verify-email, enter the emailed code and select Verify OTP. This separate page also provides Resend OTP and a return to account creation.')

page(titles[4])
h('Browse and search for a perfume')
steps(['Open Shop from the navigation.',
'Enter a product name, choose Men, Women, or Unisex when appropriate, and set a minimum or maximum price if needed.',
'Choose Default, Popularity, Price: Low to High, or Price: High to Low to organize the results.',
'Open a product to inspect its photographs, description, size, price, and available stock. Select an image thumbnail to change the main image.',
'For scent-note filtering, sign in and open Search. Combine the name, note, and category filters to narrow the results.'])
h('Save and revisit favorites')
steps(['Sign in and select the heart-shaped Favorite control on a product.',
'Open Favorites to revisit the saved collection.',
'Open a saved perfume to inspect or purchase it. Toggle its Favorite control again to remove it from the collection.'])
note('Favorites are associated with the user in the current browser’s local storage. They may not appear in another browser or after local data is cleared.')
h('Use the Adaptive Scent Forecast')
steps(['Open Home and expand the Adaptive Scent Forecast control.',
'Enable the feature and allow browser location access, or select Use Default Location.',
'Review the displayed weather conditions, forecast, and recommended perfumes. The default location is Marinduque.',
'Select a recommended product to open its details. If weather is unavailable, try the default location or browse Shop directly.'])
p('Recommendations match forecast conditions with the perfume’s scent notes. They support discovery; the customer selects the preferred fragrance.')

page(titles[5])
h('Add items and prepare the cart')
steps(['Sign in and open the required perfume. Select Add to Cart.',
'Choose a quantity within the available stock, then select Add.',
'Open Cart. Use the plus or minus controls to change quantities, or Remove to delete an unwanted item.',
'Select the checkbox for every item to be purchased. Review the estimated total and select Continue to Checkout.'])
h('Complete checkout')
steps(['In Shipping, verify Full Name and the required address fields: Barangay, City/Municipality, Province/State, ZIP Code, and Country. Add Street Address where useful.',
'Select Next. In Contact, verify Mobile Number and Email.',
'Select Next to open Review. Check the displayed information and the Cash on Delivery payment method. Use Back to correct earlier fields.',
'Select Confirm Order once and wait for processing to finish.',
'When the application opens Orders, verify the order entry, items, quantities, and amount. Record the order identifier for later inquiries.'])
h('Purchase one product directly')
steps(['On Product Detail, select Buy Now.',
'Choose the quantity and select Continue to Checkout.',
'Complete Shipping, Contact, and Review as described above. The direct purchase uses the selected product and quantity.'])
note('Checkout updates the saved profile with the supplied contact and address details. Cash on Delivery is the implemented payment method. The Delivery Notes field is displayed but is not included in the order submission; communicate essential instructions directly to the store.')

page(titles[6])
h('Track an order')
steps(['Open Orders at /orders. The current page heading is displayed as “Order Tacking.”',
'Locate the order and item identifiers. Review quantity, unit price, total, date, status, and tracking information.',
'Interpret To ship as preparation, To receive as the receiving stage, and To review as the stage that exposes review and repeat-purchase actions.',
'Revisit the page to retrieve updated records. Contact the store with the order identifier if clarification is needed.'])
h('Cancel an eligible item or buy again')
steps(['Locate the item and select its cancellation control when available.',
'Read the confirmation question. Select Yes to request cancellation, or No to retain the item.',
'Verify the resulting item status. Cancellation depends on the item’s current processing state.',
'For an item showing To Review, select Buy Again to prepare a repeat purchase, or Rate this Product to open Product Detail. Review stock and checkout details before confirming a repeat order.'])
h('Manage order notifications')
steps(['Open Notifications.',
'Use All inboxes, Unread inboxes, or Opened inboxes to filter the list. The controls display icons and entry counts.',
'Select a notification message to mark it opened.',
'Select Delete on an opened entry to hide it from the notification list.'])
note('Notification entries are derived from order records. Opened and deleted states are stored in the current browser; deleting a notification does not cancel or delete the order.')

page(titles[7])
h('Update profile, address, and appearance')
steps(['Open Account and review Profile & Contacts.',
'Edit the name, email, mobile number, and address information as required. An optional profile picture and backup address are available.',
'Select Update Profile and wait for the success confirmation.',
'Choose a Theme Preference to adjust the customer interface appearance.'])
h('Change an available account password')
steps(['In Account, open Password Manager when it is available.',
'Enter Current Password, New Password, and Confirm New Password.',
'Select Update Password and confirm the success message. If the account does not have a managed password, this section may not be shown.'])
h('Submit a product review')
steps(['Open the perfume’s Product Detail page while signed in.',
'Under Rate this scent, choose a rating from one to five stars and enter a comment.',
'Optionally attach supported media, then select Submit Review.',
'Check the Reviews section. Use Delete Review on your own review if removal is necessary.'])
h('Send store feedback')
steps(['Open Feedback and enter the relevant Order ID.',
'Choose a rating from 1 to 5 and enter a clear feedback message.',
'Optionally attach an image or supported video, then select Submit Feedback.',
'Wait for “Feedback submitted. Thank you!” before leaving the page.'])
note('The attachment control permits one image or supported video up to 15 MB; videos are limited to 30 seconds. Follow any additional validation message displayed during upload.')

page(titles[8])
h('Enter Admin Studio')
steps(['Open /admin/login.',
'Enter the authorized Admin Email and Password, then submit the sign-in form.',
'Confirm that Dashboard and the Admin Studio navigation appear. Select Menu on smaller screens when needed.'])
h('Read dashboard indicators')
p('Review Total Orders, Revenue, Low Stock, and Active SKUs to identify items requiring attention. These are application indicators: the Active SKUs value is based on the loaded inventory list, and recorded order revenue should not automatically be treated as cash already collected.')
h('Update storefront content')
steps(['Open Dashboard and locate the content block to be changed.',
'For rotating banners, supply image URLs or upload files; the bulk upload accepts up to five banners. Select Save Banners.',
'For associated banner messages, edit the relevant text and select Save Banner Content.',
'For featured content, update images using Save Featured Images and the associated titles or messages using Save Featured Pop-up Content.',
'For the main storefront image, update Hero Image URL or upload an image, then select Save Hero Image.',
'For login announcements, add or edit the available announcement entries and select Save Announcements.',
'For the post-login promotional image, edit the Pop-up Image URL or upload an image, then select Save Pop-up.',
'Wait for the save result for each block. Review the customer-facing content in a separate customer session.'])
note('Dashboard blocks have separate save buttons. Save each edited block individually. The customer promotional pop-up is normally shown once per user session, so it may not reopen immediately in the same session.')

page(titles[9])
h('Create a product listing')
steps(['Select Add Product from Admin Studio.',
'Enter Product Name, Price, Stock, Size, Notes, and Short Description. Use clear scent notes to support fragrance discovery.',
'Supply up to four Image URLs or use the upload controls. Wait until each upload completes.',
'Choose the category and status available in the form. Active maps to an active listing; Draft and Out of Stock map to an inactive listing.',
'Select Save Product and wait for “Product added.”',
'Open Modify Products to check the saved listing. Confirm its name, price, stock, images, and category.'])
note('Although Add Product displays a SKU field, the current submission does not send that value. Do not use the field as evidence that a custom SKU was saved.')
h('Modify an existing product')
steps(['Select Modify Products and locate the required listing.',
'Select its edit control to load the product information.',
'Update the name, price, stock, scent notes, size, category, description, or image URLs as required.',
'Select Save Changes and wait for “Product updated.”',
'Review the listing again and confirm the customer-facing information.'])
h('Category consistency')
p('Add Product offers Signature, Fresh, Floral, and Amber. Modify Products offers Men, Women, and Unisex, which align with the customer category filters. Review the saved category after creation and choose the appropriate customer-facing category in Modify Products when necessary.')
note('Product availability depends on saved stock and listing state. Verify the actual listing after changes rather than relying only on a selected status label.')

page(titles[10])
h('Process order items')
steps(['Open View Orders and identify the correct customer, order, and item.',
'Select To Ship, To Receive, or To Review for the item according to the actual fulfillment stage.',
'Confirm that the item’s displayed status updates. Cancelled items are shown as cancelled by the user.',
'If an order record must be removed, select Remove and read the irreversible-removal confirmation before proceeding.'])
h('Review users and feedback')
steps(['Open View Users to inspect the customer records.',
'If removal is required, verify the correct user, select Remove, and read the confirmation. The interface identifies removal as irreversible.',
'Open Feedback to review the order reference, customer details, message, rating, and any attachment. The page provides viewing; it does not expose a reply workflow.'])
h('Review and print the sales report')
steps(['Open Sales Report and allow the summary to load.',
'Review the Sales Summary, daily chart, weekly total, average daily amount, and most-sold product information.',
'Select Print PDF to open the printable report.',
'In the browser print dialog, choose the PDF destination or printer, review the preview, and save or print.',
'Select Sign Out after completing administrative work.'])
note('Sales calculations exclude cancelled and removed records in the relevant totals. They describe recorded order activity, not independently verified COD cash receipts. The chart presents a recent weekly period rather than a user-selected reporting range.')

page(titles[11])
h('Common issues and corrective actions')
icon_entries([
('OTP not received','Registration','Check the email address and spam folder. Wait for the cooldown, then select Resend OTP and use the newest code.'),
('Sign-in rejected','Login','Verify the identifier and password; accept the terms. Ask the administrator for help if the problem continues.'),
('No products found','Shop / Search','Clear the name, note, category, or price filters and try again.'),
('Checkout rejected','Checkout','Complete required fields, select at least one cart item, and verify available stock. Check Orders before repeating an uncertain submission.'),
('Upload fails','Media controls','Check file type, size, video duration, and connection. Retry after correcting the displayed error.'),
('Admin page redirects','Admin Studio','Sign in again at /admin/login using an authorized administrator account.')],(2200,1900,5260))
h('Display-only pages and operating limits')
p('Billing presents editable fields and Confirm Billing, but the control has no submission handler. Email Order Confirmation contains fixed sample details; Resend Email and Download Receipt have no action handlers. Use Checkout and Orders for actual purchases and records.')
p('Checkout text mentions SMS confirmation and a delivery PIN, but the reviewed customer workflow does not implement a complete SMS/PIN confirmation procedure. The sample confirmation page does not establish actual delivery timing or a customer’s real PIN.')
h('Documentation basis')
p('Reviewed sources include src/App.jsx for routes and access controls; src/pages/ for customer and administrator workflows; src/components/ for navigation, adaptive scent, and media controls; and server/src/routes/ for order and review operations. This manual documents the available implementation without claiming live transaction testing.','Small')

doc.core_properties.title='Severino Perfume Web Application — User Manual'
doc.core_properties.subject='Customer and administrator page reference and operating procedures'
doc.core_properties.author=''
doc.core_properties.keywords='Severino, user manual, documentation, customer, administrator'
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
with zipfile.ZipFile(OUT) as z:
    xml=z.read('word/document.xml').decode()
    assert xml.count('<w:bookmarkStart')==len(titles)
    assert xml.count('<w:hyperlink')==len(titles)
    assert len(doc.tables)==0
    assert xml.count('Primary Use:')==24
    assert 'Confirm Order' in xml and 'Table of Contents' in xml
    assert not any(x in xml for x in ['[[TOC]]','TODO','TBD'])
    for name in z.namelist():
        if name.endswith('.xml'):
            from lxml import etree
            etree.fromstring(z.read(name))
json.dump({'preset':'compact_reference_guide','cover':'editorial_cover','sections':titles,'toc':'static internal links without unverified page numbers','paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'output_bytes':OUT.stat().st_size},open(QA/'structural_audit.json','w'),indent=2)
print(OUT)
print('Structural validation passed.')
